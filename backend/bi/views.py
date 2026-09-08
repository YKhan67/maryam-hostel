# backend/bi/views.py
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.http import HttpResponse
from django.db.models import Q
from django.utils import timezone
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter
import logging

from .models import ReportTemplate, SavedReport, CustomReport, CustomReportResult
from .serializers import (
    ReportTemplateSerializer,
    SavedReportSerializer,
    CustomReportSerializer,
    CustomReportResultSerializer
)
from .services import ReportService
from .report_builder import ReportBuilder

logger = logging.getLogger(__name__)

class ReportViewSet(viewsets.ViewSet):
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return ReportTemplate.objects.filter(is_active=True)

    @action(detail=False, methods=['post'])
    def generate(self, request):
        report_type = request.data.get('report_type')
        filters = request.data.get('filters', {})

        user = request.user
        if not self._has_report_access(user, report_type):
            return Response(
                {"error": "Permission denied"},
                status=status.HTTP_403_FORBIDDEN
            )

        report_methods = {
            'REVENUE': ReportService.generate_revenue_report,
            'FEE_COLLECTION': ReportService.generate_fee_collection_report,
            'OCCUPANCY': ReportService.generate_occupancy_report,
            'INVENTORY': ReportService.generate_inventory_report,
            'PAYROLL': ReportService.generate_payroll_report,
            'MEAL_FEEDBACK': ReportService.generate_meal_feedback_report,
            'SECURITY_DEPOSIT': ReportService.generate_security_deposit_report,
            'HOSTEL_PERFORMANCE': ReportService.generate_hostel_performance_report,
            'VENDOR_ANALYSIS': ReportService.generate_vendor_analysis_report,
            'STOCK_LEVEL': ReportService.generate_stock_level_report,
        }

        generator = report_methods.get(report_type)
        if not generator:
            return Response(
                {"error": "Report type not found"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            report_data = generator(filters)
            return Response(report_data, status=status.HTTP_200_OK)
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=False, methods=['get'])
    def types(self, request):
        report_types = [
            {'id': 'REVENUE', 'name': 'Revenue Report'},
            {'id': 'FEE_COLLECTION', 'name': 'Fee Collection Report'},
            {'id': 'OCCUPANCY', 'name': 'Occupancy Report'},
            {'id': 'INVENTORY', 'name': 'Inventory Report'},
            {'id': 'PAYROLL', 'name': 'Payroll Summary'},
            {'id': 'MEAL_FEEDBACK', 'name': 'Meal Feedback Report'},
            {'id': 'SECURITY_DEPOSIT', 'name': 'Security Deposit Report'},
            {'id': 'HOSTEL_PERFORMANCE', 'name': 'Hostel Performance'},
            {'id': 'VENDOR_ANALYSIS', 'name': 'Vendor Analysis'},
            {'id': 'STOCK_LEVEL', 'name': 'Stock Level Report'},
        ]
        user = request.user
        if user.role == 'STUDENT':
            return Response([], status=status.HTTP_200_OK)
        return Response(report_types, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post'])
    def export(self, request):
        report_type = request.data.get('report_type')
        report_data = request.data.get('report_data', {})
        user = request.user

        if not self._has_report_access(user, report_type):
            return Response(
                {"error": "Permission denied"},
                status=status.HTTP_403_FORBIDDEN
            )

        try:
            wb = openpyxl.Workbook()
            ws1 = wb.active
            ws1.title = "Summary"
            ws1.append(['Report:', report_type])
            ws1.append(['Generated:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
            ws1.append([])
            summary = report_data.get('summary', {})
            for key, value in summary.items():
                ws1.append([key.replace('_', ' ').title(), value])

            ws2 = wb.create_sheet("Data")
            data_key = None
            for k in ['monthly_data', 'hostel_data', 'vendor_data', 'top_items']:
                if k in report_data:
                    data_key = k
                    break
            if data_key and report_data.get(data_key):
                data = report_data[data_key]
                if data:
                    headers = list(data[0].keys())
                    ws2.append(headers)
                    for row in data:
                        ws2.append([row.get(h, '') for h in headers])

            for ws in [ws1, ws2]:
                for column in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(column[0].column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_len:
                                max_len = len(str(cell.value))
                        except:
                            pass
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

            response = HttpResponse(
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = f'attachment; filename="{report_type}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
            wb.save(response)
            return response
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    def _has_report_access(self, user, report_type):
        if user.role in ['SUPER_ADMIN', 'CITY_MANAGER']:
            return True
        if user.role == 'PARTNER':
            return report_type in ['REVENUE', 'FEE_COLLECTION', 'HOSTEL_PERFORMANCE', 'SECURITY_DEPOSIT']
        if user.role == 'HOSTEL_MANAGER':
            return report_type in ['OCCUPANCY', 'INVENTORY', 'PAYROLL', 'MEAL_FEEDBACK', 'STOCK_LEVEL', 'VENDOR_ANALYSIS']
        if user.role == 'STAFF':
            return report_type in ['OCCUPANCY', 'INVENTORY', 'MEAL_FEEDBACK', 'STOCK_LEVEL']
        return False

class ReportTemplateViewSet(viewsets.ModelViewSet):
    queryset = ReportTemplate.objects.all()
    serializer_class = ReportTemplateSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role == 'SUPER_ADMIN':
            return super().get_queryset()
        return super().get_queryset().filter(created_by=user)

class SavedReportViewSet(viewsets.ModelViewSet):
    queryset = SavedReport.objects.all()
    serializer_class = SavedReportSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return super().get_queryset().filter(created_by=self.request.user)

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    @action(detail=True, methods=['post'])
    def run(self, request, pk=None):
        saved = self.get_object()
        report_type = saved.template.report_type
        filters = saved.filters
        generator = getattr(ReportService, f'generate_{report_type.lower()}_report', None)
        if not generator:
            return Response(
                {"error": f"Report type '{report_type}' not found"},
                status=status.HTTP_400_BAD_REQUEST
            )
        try:
            data = generator(filters)
            saved.data = data
            saved.last_run = timezone.now()
            saved.save()
            return Response(data, status=status.HTTP_200_OK)
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

class CustomReportViewSet(viewsets.ModelViewSet):
    serializer_class = CustomReportSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        return CustomReport.objects.filter(created_by=user) | CustomReport.objects.filter(is_shared=True)

    def create(self, request, *args, **kwargs):
        data = request.data.copy()
        # Set defaults for optional fields
        data.setdefault('schedule_frequency', None)
        data.setdefault('schedule_time', None)
        data.setdefault('schedule_recipients', [])
        data.setdefault('is_scheduled', False)
        data.setdefault('distinct', False)
        data.setdefault('tables', [])
        data.setdefault('fields', [])
        data.setdefault('filters', [])
        data.setdefault('group_by', [])
        data.setdefault('sort_by', [])
        data.setdefault('date_range', {})
        data.setdefault('is_shared', False)
        data.setdefault('shared_with_roles', [])
        
        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        data = request.data.copy()
        
        # Preserve existing values if not provided
        data.setdefault('schedule_frequency', instance.schedule_frequency)
        data.setdefault('schedule_time', instance.schedule_time)
        data.setdefault('schedule_recipients', instance.schedule_recipients)
        data.setdefault('is_scheduled', instance.is_scheduled)
        data.setdefault('distinct', instance.distinct)
        data.setdefault('tables', instance.tables)
        data.setdefault('fields', instance.fields)
        data.setdefault('filters', instance.filters)
        data.setdefault('group_by', instance.group_by)
        data.setdefault('sort_by', instance.sort_by)
        data.setdefault('date_range', instance.date_range)
        data.setdefault('is_shared', instance.is_shared)
        data.setdefault('shared_with_roles', instance.shared_with_roles)
        
        serializer = self.get_serializer(instance, data=data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response(serializer.data)

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    @action(detail=False, methods=['get'])
    def tables(self, request):
        try:
            tables = ReportBuilder.get_available_tables()
            return Response(tables, status=status.HTTP_200_OK)
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=False, methods=['post'])
    def generate(self, request):
        try:
            tables = request.data.get('tables', [])
            fields = request.data.get('fields', [])
            filters = request.data.get('filters', [])
            group_by = request.data.get('group_by', [])
            sort_by = request.data.get('sort_by', [])
            date_range = request.data.get('date_range', {})
            distinct = request.data.get('distinct', False)
            report_id = request.data.get('report_id')

            print(f"Generate: tables={tables}, fields={len(fields)}, filters={filters}")

            if not tables or not fields:
                return Response(
                    {"error": "Tables and fields are required"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # `distinct` was previously read but never passed through, so the
            # "DISTINCT" checkbox in the report builder had no effect. `fields`
            # is now also passed to execute_query so result keys are derived
            # from the actual field names instead of a lossy alias parse.
            query = ReportBuilder.build_query(tables, fields, filters, group_by, sort_by, date_range, distinct)

            # build_query() returns "" both for invalid tables/fields and when
            # no join path (even indirect, multi-hop) exists between the
            # selected tables. Previously an unrelated-tables case fell
            # through to CROSS JOIN and silently produced wrong data instead
            # of an error - surface it clearly here instead.
            if not query:
                return Response(
                    {"error": "Couldn't find a relationship between the selected tables. Try adding a table that links them (e.g. a student/profile table between a user table and a fee/record table)."},
                    status=status.HTTP_400_BAD_REQUEST
                )

            results = ReportBuilder.execute_query(query, fields)

            for row in results:
                for k, v in row.items():
                    if isinstance(v, bool):
                        row[k] = 1 if v else 0

            if report_id:
                try:
                    report = CustomReport.objects.get(id=report_id, created_by=request.user)
                    CustomReportResult.objects.create(
                        report=report,
                        data={'results': results},
                        generated_by=request.user
                    )
                    report.last_run = timezone.now()
                    report.save()
                except CustomReport.DoesNotExist:
                    pass

            return Response({
                "results": results,
                "count": len(results),
                "query": query
            }, status=status.HTTP_200_OK)

        except Exception as e:
            print(f"Generate error: {e}")
            import traceback
            traceback.print_exc()
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=False, methods=['post'])
    def save_report(self, request):
        """Save or update a custom report configuration."""
        try:
            print(f"Save report request data: {request.data}")

            name = request.data.get('name', '').strip()
            if not name:
                return Response(
                    {"error": "Report name is required"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Build the report data, handling the 'distinct' field
            report_data = {
                'name': name,
                'description': request.data.get('description', ''),
                'tables': request.data.get('tables', []),
                'fields': request.data.get('fields', []),
                'filters': request.data.get('filters', []),
                'group_by': request.data.get('group_by', []),
                'sort_by': request.data.get('sort_by', []),
                'date_range': request.data.get('date_range', {}),
                'distinct': request.data.get('distinct', False),
                'is_shared': request.data.get('is_shared', False),
                'shared_with_roles': request.data.get('shared_with_roles', []),
                'is_scheduled': request.data.get('is_scheduled', False),
                'schedule_frequency': request.data.get('schedule_frequency', None),
                'schedule_time': request.data.get('schedule_time', None),
                'schedule_recipients': request.data.get('schedule_recipients', []),
            }

            if not report_data['tables']:
                return Response(
                    {"error": "At least one table is required"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            if not report_data['fields']:
                return Response(
                    {"error": "At least one field is required"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            report_id = request.data.get('id')
            if report_id:
                try:
                    report = CustomReport.objects.get(id=report_id, created_by=request.user)
                    for key, value in report_data.items():
                        setattr(report, key, value)
                    report.save()
                    serializer = CustomReportSerializer(report)
                    return Response(serializer.data, status=status.HTTP_200_OK)
                except CustomReport.DoesNotExist:
                    return Response(
                        {"error": "Report not found or not owned"},
                        status=status.HTTP_404_NOT_FOUND
                    )
            else:
                # Create new report
                serializer = CustomReportSerializer(data=report_data)
                if serializer.is_valid():
                    report = serializer.save(created_by=request.user)
                    return Response(serializer.data, status=status.HTTP_201_CREATED)
                else:
                    # Log validation errors for debugging
                    logger.error(f"Validation errors: {serializer.errors}")
                    return Response(
                        {"error": "Validation failed", "details": serializer.errors},
                        status=status.HTTP_400_BAD_REQUEST
                    )

        except Exception as e:
            logger.error(f"Save error: {e}")
            import traceback
            traceback.print_exc()
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=True, methods=['get'])
    def results(self, request, pk=None):
        try:
            report = self.get_object()
            results = CustomReportResult.objects.filter(report=report).order_by('-generated_at')
            serializer = CustomReportResultSerializer(results, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=True, methods=['post'])
    def duplicate(self, request, pk=None):
        try:
            original = self.get_object()
            new_report = CustomReport(
                name=f"{original.name} (Copy)",
                description=original.description,
                tables=original.tables,
                fields=original.fields,
                filters=original.filters,
                group_by=original.group_by,
                sort_by=original.sort_by,
                date_range=original.date_range,
                distinct=original.distinct,
                is_shared=original.is_shared,
                shared_with_roles=original.shared_with_roles,
                created_by=request.user
            )
            new_report.save()
            serializer = CustomReportSerializer(new_report)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        except Exception as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )