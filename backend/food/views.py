# backend/food/views.py

import io
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime, date, timedelta
from django.http import HttpResponse
from django.db.models import Sum, Q
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated

from .models import (
    MealCategory, Meal, MealRecipe, DailyMenu,
    WeeklyMenuTemplate, MealFeedback, GroceryRequirement
)
from .serializers import (
    MealCategorySerializer, MealSerializer, MealRecipeSerializer,
    DailyMenuSerializer, WeeklyMenuTemplateSerializer,
    MealFeedbackSerializer, GroceryRequirementSerializer
)
from hostels.models import Hostel, StudentProfile
from inventory.models import Item


class MealCategoryViewSet(viewsets.ModelViewSet):
    queryset = MealCategory.objects.filter(is_active=True)
    serializer_class = MealCategorySerializer
    permission_classes = [permissions.IsAuthenticated]


class MealViewSet(viewsets.ModelViewSet):
    queryset = Meal.objects.filter(is_active=True)
    serializer_class = MealSerializer
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=True, methods=['get', 'post'])
    def recipes(self, request, pk=None):
        meal = self.get_object()

        if request.method == 'GET':
            recipes = meal.recipes.all()
            serializer = MealRecipeSerializer(recipes, many=True)
            return Response(serializer.data)

        serializer = MealRecipeSerializer(
            data=request.data,
            context={'request': request}
        )

        if serializer.is_valid():
            serializer.save(meal=meal)
            return Response(
                serializer.data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    @action(
        detail=True,
        methods=['patch', 'delete'],
        url_path=r'recipes/(?P<recipe_id>[^/.]+)',
        url_name='recipe-detail'
    )
    def recipe_detail(self, request, pk=None, recipe_id=None):
        meal = self.get_object()

        try:
            recipe = meal.recipes.get(pk=recipe_id)
        except MealRecipe.DoesNotExist:
            return Response(
                {"error": "Recipe ingredient not found for this meal."},
                status=status.HTTP_404_NOT_FOUND
            )

        if request.method == 'PATCH':
            serializer = MealRecipeSerializer(
                recipe,
                data=request.data,
                partial=True,
                context={'request': request}
            )

            if serializer.is_valid():
                serializer.save()
                return Response(serializer.data)

            return Response(
                serializer.errors,
                status=status.HTTP_400_BAD_REQUEST
            )

        recipe.delete()

        return Response(status=status.HTTP_204_NO_CONTENT)


class DailyMenuViewSet(viewsets.ModelViewSet):
    queryset = DailyMenu.objects.all()
    serializer_class = DailyMenuSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        queryset = super().get_queryset()

        if user.role == "SUPER_ADMIN":
            return queryset.order_by('-date', 'meal_type')

        if hasattr(user, 'role') and user.role == 'STUDENT':
            try:
                student = user.student_profile
                queryset = queryset.filter(hostel=student.hostel)
            except StudentProfile.DoesNotExist:
                return queryset.none()
        elif hasattr(user, 'role') and user.role in ['HOSTEL_MANAGER', 'PARTNER', 'STAFF'] and user.hostel:
            queryset = queryset.filter(hostel=user.hostel)
        else:
            return queryset.none()

        return queryset.order_by('-date', 'meal_type')

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'])
    def today(self, request):
        today = date.today()
        user = request.user

        if user.role == "SUPER_ADMIN":
            menu = DailyMenu.objects.filter(date=today).order_by('meal_type')
            serializer = DailyMenuSerializer(menu, many=True)
            return Response(serializer.data)

        if hasattr(user, 'role') and user.role == 'STUDENT':
            try:
                student = user.student_profile
                hostel = student.hostel
            except StudentProfile.DoesNotExist:
                return Response({"error": "Student profile not found"}, status=400)
        elif user.hostel:
            hostel = user.hostel
        else:
            return Response([], status=200)

        menu = DailyMenu.objects.filter(hostel=hostel, date=today).order_by('meal_type')
        serializer = DailyMenuSerializer(menu, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'])
    def week(self, request):
        today = date.today()
        start_of_week = today - timedelta(days=today.weekday())
        end_of_week = start_of_week + timedelta(days=6)

        user = request.user

        if user.role == "SUPER_ADMIN":
            menu = DailyMenu.objects.filter(
                date__gte=start_of_week,
                date__lte=end_of_week
            ).order_by('date', 'meal_type')
            serializer = DailyMenuSerializer(menu, many=True)
            return Response(serializer.data)

        if hasattr(user, 'role') and user.role == 'STUDENT':
            try:
                student = user.student_profile
                hostel = student.hostel
            except StudentProfile.DoesNotExist:
                return Response({"error": "Student profile not found"}, status=400)
        elif user.hostel:
            hostel = user.hostel
        else:
            return Response([], status=200)

        menu = DailyMenu.objects.filter(
            hostel=hostel,
            date__gte=start_of_week,
            date__lte=end_of_week
        ).order_by('date', 'meal_type')

        serializer = DailyMenuSerializer(menu, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def feedback(self, request, pk=None):
        daily_menu = self.get_object()
        user = request.user

        if user.role != 'STUDENT':
            return Response({"error": "Only students can provide feedback"}, status=403)

        try:
            student = user.student_profile
        except StudentProfile.DoesNotExist:
            return Response({"error": "Student profile not found"}, status=400)

        rating = request.data.get('rating')
        comment = request.data.get('comment', '')

        if not rating or rating < 1 or rating > 5:
            return Response({"error": "Rating must be between 1 and 5"}, status=400)

        feedback, created = MealFeedback.objects.get_or_create(
            student=student,
            daily_menu=daily_menu,
            defaults={'rating': rating, 'comment': comment}
        )

        if not created:
            feedback.rating = rating
            feedback.comment = comment
            feedback.save()

        serializer = MealFeedbackSerializer(feedback)
        return Response(serializer.data)


class WeeklyMenuTemplateViewSet(viewsets.ModelViewSet):
    queryset = WeeklyMenuTemplate.objects.all()
    serializer_class = WeeklyMenuTemplateSerializer
    permission_classes = [permissions.IsAuthenticated]


class MealFeedbackViewSet(viewsets.ModelViewSet):
    queryset = MealFeedback.objects.all()
    serializer_class = MealFeedbackSerializer
    permission_classes = [permissions.IsAuthenticated]


class GroceryRequirementViewSet(viewsets.ModelViewSet):
    queryset = GroceryRequirement.objects.all()
    serializer_class = GroceryRequirementSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        queryset = super().get_queryset()

        if user.role == "SUPER_ADMIN":
            return queryset

        if hasattr(user, 'role') and user.role in ['HOSTEL_MANAGER', 'PARTNER', 'STAFF'] and user.hostel:
            queryset = queryset.filter(hostel=user.hostel)

        return queryset

    @action(detail=False, methods=['post'])
    def generate(self, request):
        user = request.user

        if user.role not in ['SUPER_ADMIN', 'HOSTEL_MANAGER', 'PARTNER']:
            return Response({"error": "Permission denied"}, status=403)

        hostel_id = request.data.get('hostel')
        start_date = request.data.get('start_date')
        end_date = request.data.get('end_date')
        servings_per_person = int(request.data.get('servings_per_person', 1))
        global_waste_factor = float(request.data.get('waste_factor', 0.015))

        if not hostel_id or not start_date or not end_date:
            return Response({"error": "hostel, start_date, and end_date are required"}, status=400)

        try:
            hostel = Hostel.objects.get(id=hostel_id)
        except Hostel.DoesNotExist:
            return Response({"error": "Hostel not found"}, status=404)

        active_students = StudentProfile.objects.filter(
            hostel=hostel,
            is_active=True,
            user__is_active=True
        ).count()

        if active_students == 0:
            return Response({
                "error": "No active students found in this hostel",
                "hostel": hostel.name,
                "active_students": 0
            }, status=400)

        menus = DailyMenu.objects.filter(
            hostel=hostel,
            date__gte=start_date,
            date__lte=end_date
        ).select_related('meal')

        if not menus.exists():
            return Response({
                "error": "No meals scheduled in this date range",
                "hostel": hostel.name,
                "active_students": active_students,
                "start_date": start_date,
                "end_date": end_date
            }, status=400)

        ingredient_totals = {}
        total_meals_count = menus.count()

        for daily_menu in menus:
            for recipe in daily_menu.meal.recipes.all():
                item_key = recipe.item.id
                if item_key not in ingredient_totals:
                    ingredient_totals[item_key] = {
                        'item': recipe.item,
                        'quantity': 0,
                        'unit': recipe.unit,
                        'estimated_cost': 0,
                        'waste_factor': float(recipe.waste_factor or global_waste_factor),
                    }

                base_quantity = float(recipe.quantity_required) * active_students * servings_per_person
                waste = float(recipe.waste_factor or global_waste_factor)
                quantity_with_waste = base_quantity * (1 + waste)

                ingredient_totals[item_key]['quantity'] += quantity_with_waste
                ingredient_totals[item_key]['estimated_cost'] += float(recipe.estimated_cost or 0) * quantity_with_waste

        created_items = []
        for item_id, data in ingredient_totals.items():
            grocery_item, created = GroceryRequirement.objects.get_or_create(
                hostel=hostel,
                start_date=start_date,
                end_date=end_date,
                item=data['item'],
                defaults={
                    'quantity_needed': data['quantity'],
                    'unit': data['unit'],
                    'estimated_cost': data['estimated_cost'],
                    'status': 'PENDING',
                    'notes': f"Based on {active_students} active students, {total_meals_count} meals, {servings_per_person} serving(s) per person. Waste factor: {data['waste_factor']*100:.1f}%"
                }
            )
            if created:
                created_items.append(grocery_item.id)
            else:
                grocery_item.quantity_needed = data['quantity']
                grocery_item.estimated_cost = data['estimated_cost']
                grocery_item.notes = f"Based on {active_students} active students, {total_meals_count} meals, {servings_per_person} serving(s) per person. Waste factor: {data['waste_factor']*100:.1f}%"
                grocery_item.save()

        return Response({
            "message": f"Generated grocery requirements for {len(created_items)} items",
            "items_created": len(created_items),
            "hostel": hostel.name,
            "active_students": active_students,
            "total_meals": total_meals_count,
            "servings_per_person": servings_per_person,
            "global_waste_factor": global_waste_factor,
            "date_range": f"{start_date} to {end_date}"
        })


class ExportMealTemplateView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        wb = openpyxl.Workbook()

        ws1 = wb.active
        ws1.title = "Categories"
        headers = ["Name", "Emoji", "Description"]
        ws1.append(headers)

        categories = MealCategory.objects.all()
        if categories.exists():
            for cat in categories:
                ws1.append([cat.name, cat.emoji, cat.description])
        else:
            sample_categories = [
                ["BREAKFAST", "🌅", "Morning meals"],
                ["BRUNCH", "🌤️", "Late morning meals"],
                ["LUNCH", "☀️", "Midday meals"],
                ["SNACKS", "🍪", "Between meals"],
                ["DINNER", "🌙", "Evening meals"],
            ]
            for row in sample_categories:
                ws1.append(row)

        ws2 = wb.create_sheet("Meals")
        headers = ["Name", "Description", "Category", "Dietary Tags", "Prep Time (mins)"]
        ws2.append(headers)

        meals = Meal.objects.select_related('category').all()
        if meals.exists():
            for meal in meals:
                ws2.append([
                    meal.name,
                    meal.description,
                    meal.category.name,
                    meal.dietary_tags,
                    meal.preparation_time
                ])
        else:
            sample_meals = [
                ["Paratha & Chai", "Fresh homemade paratha with masala chai", "BREAKFAST", "Vegetarian", 30],
                ["Bread & Omelette", "Toasted bread with fluffy omelette", "BREAKFAST", "Vegetarian, High Protein", 20],
                ["Chicken Biryani", "Aromatic basmati rice with chicken", "LUNCH", "Halal", 60],
                ["Rice & Curry", "Steamed rice with vegetable curry", "LUNCH", "Vegetarian, Halal", 45],
                ["Chicken Karahi", "Spicy chicken karahi with naan", "DINNER", "Halal", 45],
                ["Pasta", "Creamy pasta with vegetables", "DINNER", "Vegetarian", 35],
            ]
            for row in sample_meals:
                ws2.append(row)

        ws3 = wb.create_sheet("Schedule")
        headers = ["Date", "Hostel Name", "Meal Type", "Meal Name", "Is Featured", "Special Note"]
        ws3.append(headers)

        today = datetime.now().date()
        hostels = Hostel.objects.all()
        hostel_name = hostels[0].name if hostels.exists() else "Your Hostel Name"

        for i in range(7):
            date_obj = today + timedelta(days=i)
            ws3.append([
                date_obj.strftime("%Y-%m-%d"),
                hostel_name,
                "BREAKFAST",
                "",
                "No",
                "",
            ])
            ws3.append([
                date_obj.strftime("%Y-%m-%d"),
                hostel_name,
                "LUNCH",
                "",
                "No",
                "",
            ])
            ws3.append([
                date_obj.strftime("%Y-%m-%d"),
                hostel_name,
                "DINNER",
                "",
                "No",
                "",
            ])

        for ws in [ws1, ws2, ws3]:
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_length = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_length

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="meal_menu_template.xlsx"'
        wb.save(response)
        return response


class ImportMealExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if 'file' not in request.FILES:
            return Response(
                {"error": "No file provided"},
                status=status.HTTP_400_BAD_REQUEST
            )

        file = request.FILES['file']

        if not file.name.endswith(('.xlsx', '.xls')):
            return Response(
                {"error": "Invalid file format. Please upload .xlsx or .xls file"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            wb = openpyxl.load_workbook(file)
            results = {
                "categories_created": 0,
                "categories_updated": 0,
                "meals_created": 0,
                "meals_updated": 0,
                "schedules_created": 0,
                "schedules_updated": 0,
                "errors": [],
                "debug": []
            }

            available_hostels = list(Hostel.objects.values_list('name', flat=True))
            available_meals = list(Meal.objects.values_list('name', flat=True))
            results["debug"].append(f"Available hostels: {available_hostels}")
            results["debug"].append(f"Available meals: {available_meals}")

            if "Categories" in wb.sheetnames:
                ws = wb["Categories"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    try:
                        name = str(row[0]).strip() if row[0] else ""
                        emoji = str(row[1]).strip() if row[1] else ""
                        description = str(row[2]).strip() if row[2] else ""

                        if name:
                            category, created = MealCategory.objects.update_or_create(
                                name=name.upper(),
                                defaults={"emoji": emoji, "description": description}
                            )
                            if created:
                                results["categories_created"] += 1
                            else:
                                results["categories_updated"] += 1
                    except Exception as e:
                        results["errors"].append(f"Category error: {str(e)}")

            if "Meals" in wb.sheetnames:
                ws = wb["Meals"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    try:
                        name = str(row[0]).strip() if row[0] else ""
                        description = str(row[1]).strip() if row[1] else ""
                        category_name = str(row[2]).strip() if row[2] else ""
                        dietary_tags = str(row[3]).strip() if row[3] else ""
                        prep_time = int(row[4]) if row[4] else 30

                        if name and category_name:
                            category = MealCategory.objects.filter(name=category_name.upper()).first()
                            if not category:
                                results["errors"].append(f"Category '{category_name}' not found for meal '{name}'")
                                continue

                            meal, created = Meal.objects.update_or_create(
                                name=name,
                                defaults={
                                    "description": description,
                                    "category": category,
                                    "dietary_tags": dietary_tags,
                                    "preparation_time": prep_time
                                }
                            )
                            if created:
                                results["meals_created"] += 1
                            else:
                                results["meals_updated"] += 1
                    except Exception as e:
                        results["errors"].append(f"Meal error: {str(e)}")

            if "Schedule" in wb.sheetnames:
                ws = wb["Schedule"]
                results["debug"].append(f"Schedule sheet found. Rows count: {ws.max_row - 1}")

                for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        continue
                    try:
                        date_str = str(row[0]).strip() if row[0] else ""
                        hostel_name = str(row[1]).strip() if row[1] else ""
                        meal_type = str(row[2]).strip() if row[2] else ""
                        meal_name = str(row[3]).strip() if row[3] else ""
                        is_featured = str(row[4]).strip() if row[4] else "No"
                        special_note = str(row[5]).strip() if row[5] else ""

                        if not date_str or not hostel_name or not meal_type or not meal_name:
                            results["errors"].append(f"Row {row_idx}: Missing required fields")
                            continue

                        try:
                            date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
                        except ValueError:
                            results["errors"].append(f"Row {row_idx}: Invalid date format '{date_str}'")
                            continue

                        hostel = Hostel.objects.filter(name__icontains=hostel_name).first()
                        if not hostel:
                            results["errors"].append(f"Row {row_idx}: Hostel '{hostel_name}' not found")
                            continue

                        meal = Meal.objects.filter(name__icontains=meal_name).first()
                        if not meal:
                            results["errors"].append(f"Row {row_idx}: Meal '{meal_name}' not found")
                            continue

                        valid_meal_types = ['BREAKFAST', 'BRUNCH', 'LUNCH', 'SNACKS', 'DINNER']
                        if meal_type.upper() not in valid_meal_types:
                            results["errors"].append(f"Row {row_idx}: Invalid meal type '{meal_type}'")
                            continue

                        menu, created = DailyMenu.objects.update_or_create(
                            hostel=hostel,
                            date=date_obj,
                            meal_type=meal_type.upper(),
                            defaults={
                                "meal": meal,
                                "is_featured": is_featured.lower() in ["yes", "true", "1"],
                                "special_note": special_note
                            }
                        )
                        if created:
                            results["schedules_created"] += 1
                        else:
                            results["schedules_updated"] += 1

                    except Exception as e:
                        results["errors"].append(f"Row {row_idx}: {str(e)}")

            return Response(results, status=status.HTTP_200_OK)

        except Exception as e:
            return Response(
                {"error": f"Failed to process file: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


class ExportRecipeTemplateView(APIView):
    """
    Export Excel template for meal recipes (ingredients)
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        wb = openpyxl.Workbook()

        ws = wb.active
        ws.title = "Recipes"
        headers = ["Meal Name", "Item Name", "Quantity", "Unit", "Estimated Cost", "Waste Factor (%)"]
        ws.append(headers)

        sample_recipes = [
            ["Paratha & Chai", "Flour", "0.5", "kg", "150", "1.5"],
            ["Paratha & Chai", "Eggs", "4", "pieces", "40", "1.0"],
            ["Paratha & Chai", "Milk", "0.25", "liters", "30", "1.0"],
            ["Chicken Biryani", "Chicken", "1.5", "kg", "600", "2.0"],
            ["Chicken Biryani", "Rice", "1", "kg", "200", "1.5"],
            ["Chicken Biryani", "Onion", "2", "pieces", "30", "2.0"],
            ["Chicken Karahi", "Chicken", "1", "kg", "400", "2.0"],
            ["Chicken Karahi", "Tomato", "3", "pieces", "25", "2.5"],
            ["Chicken Karahi", "Ginger Garlic", "0.05", "kg", "15", "1.0"],
        ]
        for row in sample_recipes:
            ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_length = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_length

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="meal_recipe_template.xlsx"'
        wb.save(response)
        return response


class ImportRecipeExcelView(APIView):
    """
    Import meal recipes from Excel
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if 'file' not in request.FILES:
            return Response(
                {"error": "No file provided"},
                status=status.HTTP_400_BAD_REQUEST
            )

        file = request.FILES['file']

        if not file.name.endswith(('.xlsx', '.xls')):
            return Response(
                {"error": "Invalid file format. Please upload .xlsx or .xls file"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            wb = openpyxl.load_workbook(file)
            results = {
                "created": 0,
                "updated": 0,
                "errors": []
            }

            if "Recipes" not in wb.sheetnames:
                return Response(
                    {"error": "Sheet 'Recipes' not found in the Excel file"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            ws = wb["Recipes"]

            meal_id = request.data.get('meal_id')
            target_meal = None
            if meal_id:
                try:
                    target_meal = Meal.objects.get(id=meal_id)
                except Meal.DoesNotExist:
                    results["errors"].append(f"Meal with ID {meal_id} not found")
                    return Response(results, status=status.HTTP_400_BAD_REQUEST)

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not row[0]:
                    continue
                try:
                    meal_name = str(row[0]).strip() if row[0] else ""
                    item_name = str(row[1]).strip() if row[1] else ""
                    quantity = float(row[2]) if row[2] else 0
                    unit = str(row[3]).strip() if row[3] else "kg"
                    estimated_cost = float(row[4]) if row[4] else 0
                    waste_factor = float(row[5]) / 100 if row[5] else 0.015

                    if not meal_name or not item_name:
                        results["errors"].append(f"Row {row_idx}: Meal name and Item name are required")
                        continue

                    if target_meal:
                        meal = target_meal
                    else:
                        meal = Meal.objects.filter(name__icontains=meal_name).first()
                        if not meal:
                            results["errors"].append(f"Row {row_idx}: Meal '{meal_name}' not found. Please create it first or select a meal.")
                            continue

                    item = Item.objects.filter(name__icontains=item_name).first()
                    if not item:
                        results["errors"].append(f"Row {row_idx}: Item '{item_name}' not found. Please add it to inventory first.")
                        continue

                    recipe, created = MealRecipe.objects.update_or_create(
                        meal=meal,
                        item=item,
                        defaults={
                            'quantity_required': quantity,
                            'unit': unit,
                            'estimated_cost': estimated_cost,
                            'waste_factor': waste_factor,
                        }
                    )
                    if created:
                        results["created"] += 1
                    else:
                        results["updated"] += 1

                except Exception as e:
                    results["errors"].append(f"Row {row_idx}: {str(e)}")

            return Response(results, status=status.HTTP_200_OK)

        except Exception as e:
            return Response(
                {"error": f"Failed to process file: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )